"""Generate Part 1 personal report (.docx) for CW2 Group 12."""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = r"C:\Users\gjk20\Desktop\personal report.docx"

doc = Document()

section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21.0)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def add_heading(text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    if level == 0:
        run.font.size = Pt(16)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 1:
        run.font.size = Pt(13)
    else:
        run.font.size = Pt(11)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)


def add_para(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    # support **bold** segments
    parts = text.split("**")
    for i, seg in enumerate(parts):
        run = p.add_run(seg)
        if i % 2 == 1:
            run.bold = True
    return p


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    parts = text.split("**")
    for i, seg in enumerate(parts):
        run = p.add_run(seg)
        if i % 2 == 1:
            run.bold = True
    return p


# -------- Title --------
add_heading("CW2 Personal Report (Part 1: Individual Reflection)", level=0)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
mr = meta.add_run(
    "Group 12 — E-Scooter Rental System    |    Role: Backend (Spring Boot) & Frontend (uni-app x) Developer"
)
mr.italic = True
mr.font.size = Pt(10)

# -------- 1.1 --------
add_heading("1.1 Your role in the team (5 marks)", level=1)

add_para(
    "At the start of the project, the team carried out a Belbin Self-Perception Inventory. "
    "My profile came back with two clearly dominant roles: **Implementer / Company Worker (CW)** "
    "as the primary role, with **Completer-Finisher (CF)** as a strong secondary. I therefore "
    "established myself in the team as the Implementer — the person who would convert agreed "
    "plans, user stories and UI mock-ups into running code on both the Spring Boot backend and "
    "the uni-app x frontend, and who would polish loose ends before every sprint demo."
)

add_para(
    "Reflecting back on that Belbin role, I believe I performed it consistently and to a level "
    "that the rest of the team relied on. Two specific pieces of work make this concrete."
)

add_para(
    "**Example 1 — End-to-end booking and simulated card payment flow (Sprint 2 → Sprint 3).** "
    "Once the team had signed off the booking workflow, I implemented the full path: "
    "BookingController → BookingServiceImpl → JPA repository on the backend, plus the matching "
    "booking, payment and confirmation pages on the frontend (pages/escooters, pages/payment, "
    "pages/orders). The booking-to-payment integration is captured in the commit "
    "\"feat: add booking payment flow and refresh scooter app\". A few days later, the tester "
    "found that the price could be tampered with from the client side, and I addressed it in "
    "\"fix: payment amount now uses server-calculated price; improve booking security\" by moving "
    "the price calculation into BillingServiceImpl and validating the server-side amount before "
    "the simulated card step. This is exactly the Implementer behaviour Belbin describes — "
    "turning plans into systematic, deliverable actions, including the unglamorous parts (input "
    "validation, error DTOs, edge cases) that are easy to skip."
)

add_para(
    "**Example 2 — Concurrency-safe bookings, admin dashboard and analytics (Sprint 3).** "
    "When two clients tried to book the same scooter almost simultaneously, the original "
    "implementation allowed double-bookings. I introduced a status-guarded conditional update "
    "plus a scooter cache control layer (commit \"improve booking concurrency and scooter cache "
    "control\"), which directly satisfies product backlog ID 23 (multi-client concurrent usage). "
    "I then delivered the admin-side analytics — weekly income per duration bucket and combined "
    "daily income across a week (ID 19, ID 20, ID 21) — by wiring RevenueStatsDTO, "
    "DurationRevenueDTO and DailyTrendPointDTO into the admin dashboard charts (commits "
    "\"feat: upgrade admin dashboard UI and add data analytics charts\", \"feat: improve admin "
    "billing and fleet management\", \"feat: modernize admin ui system\"). My secondary "
    "Completer-Finisher trait also showed up here: before each demo I ran the cleanup commits "
    "(e.g. \"refactor: unify backend errors and frontend messages\", \"chore: stop tracking "
    "frontend node_modules\") so the demo build was stable and the codebase did not accumulate "
    "debt between sprints."
)

# -------- 1.2 --------
add_heading("1.2 Your strengths as part of the team (2 marks)", level=1)

add_para(
    "Looking at the overall Belbin profile of Group 12, the team was strong in Resource "
    "Investigator and Shaper roles but had only one developer who was genuinely comfortable on "
    "both sides of the API boundary. Against that profile, my particular strength was "
    "**delivering vertically integrated, full-stack increments quickly and with a stable "
    "contract between the backend and the frontend.**"
)

add_para(
    "A concrete example is Sprint 3, where the team needed to deliver product backlog ID 9 "
    "(staff taking bookings for unregistered users) and ID 23 (multi-client concurrent usage) "
    "in the same week. I designed the backend endpoints — the guest-booking flow and the "
    "admin-side guest booking controller — and, in the same branch, implemented the matching "
    "pages/guest-booking and pages/admin screens. Because both sides were authored together, "
    "the request/response contract was consistent on the first integration attempt, captured in "
    "the commit \"realize ID9 and ID23\". The benefit to the team was tangible: the teammate "
    "responsible for system testing was able to start writing test scripts against a stable "
    "schema almost immediately; the Resource Investigator did not have to re-edit the demo "
    "slides because of late field renames; and the Sprint 3 demo build was frozen two days "
    "ahead of the deadline, which gave us time to rehearse the presentation rather than "
    "firefight integration bugs."
)

# -------- 1.3 --------
add_heading("1.3 Your weaknesses as part of the team (3 marks)", level=1)

add_para(
    "The flip side of a strong Implementer is the weakness Belbin explicitly associates with "
    "the role: **a degree of inflexibility and being slow to respond to new possibilities.** "
    "In practice for me this took the form of pulling too much work onto myself and "
    "under-communicating early design decisions, which forced teammates to react to changes "
    "rather than shape them."
)

add_para(
    "A specific example occurred in Sprint 2. Over one weekend I implemented the early version "
    "of the booking and payment flow almost end-to-end on my own, without first agreeing the "
    "DTO field names with the teammate writing the system test scripts. When the test scripts "
    "were later run against my code, several cases failed because I had silently changed field "
    "names (for instance between bookingId / id, and between priceCents and amount). The impact "
    "on the team was twofold: the tester had to rewrite roughly nine system tests, costing "
    "about half a day; and the Sprint 2 status slides had to be updated at the last minute "
    "because they referenced the old names. It also briefly damaged trust — the team could no "
    "longer assume an interface was stable just because I had \"finished\" it."
)

add_para("I dealt with this weakness in three concrete steps before Sprint 3:")

add_bullet(
    "I refactored the backend to expose stable response DTOs (BookingResponse, ScooterResponse, "
    "FeedbackResponse, RevenueStatsDTO and similar) and committed \"refactor: unify backend "
    "errors and frontend messages\", so error codes and field names became a single source of "
    "truth across the codebase."
)
add_bullet(
    "I added an OpenAPI / Swagger configuration (OpenApiConfig) so teammates could read the "
    "live API contract themselves rather than waiting on me to explain it in chat."
)
add_bullet(
    "I changed my personal workflow: I now open a \"thin slice\" PR with the DTO and an empty "
    "service first, ask the team to react in GitHub comments, and only then write the body. "
    "This deliberately forces me out of the head-down Implementer pattern."
)

add_para(
    "By Sprint 3 the effect was measurable. When I implemented product backlog items ID 14, "
    "ID 15, ID 20 and ID 25 (commit \"realize ID14 15 20 and 25\"), the DTOs had been reviewed "
    "in advance, the system test scripts ran against my branch with no rework, and the tester "
    "explicitly noted in the team meeting that integration was \"smoother than last sprint\". "
    "I still recognise the Implementer weakness in myself — under deadline pressure my instinct "
    "is still to write the code first and discuss it later — but I now have habits and tooling "
    "(Swagger, thin-slice PRs, shared error codes) that compensate for it inside the team."
)

doc.save(OUTPUT)
print(f"Saved: {OUTPUT}")
